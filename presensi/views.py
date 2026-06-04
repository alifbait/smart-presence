from django.shortcuts import render, redirect
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.utils import timezone
from datetime import time
from karyawan.models import Pegawai
from presensi.models import Presensi, PengajuanIzin
from presensi.forms import PresensiForm, PengajuanIzinForm

import pandas as pd
from django.http import HttpResponse
from io import BytesIO

# Reportlab imports
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfgen import canvas

def get_pegawai_atau_salah(request):
    """
    Helper untuk mendapatkan data Pegawai dari user login secara aman.
    Mengembalikan objek Pegawai jika ada, atau None jika tidak terhubung.
    """
    if hasattr(request.user, 'pegawai'):
        return request.user.pegawai
    return None


@login_required
def presensi_harian(request):
    """
    Halaman utama absensi. Menampilkan waktu digital dan tombol absensi masuk/pulang.
    """
    pegawai = get_pegawai_atau_salah(request)
    if not pegawai:
        return render(request, 'presensi/no_profile.html', {'judul_halaman': 'Akses Dibatasi | Smart Presence'})

    now = timezone.localtime(timezone.now())
    today = now.date()
    
    # Ambil data presensi hari ini jika sudah ada
    presensi_hari_ini = Presensi.objects.filter(pegawai=pegawai, tanggal=today).first()
    form = PresensiForm(instance=presensi_hari_ini)

    konteks = {
        'judul_halaman': 'Absensi Harian | Smart Presence',
        'pegawai': pegawai,
        'presensi_hari_ini': presensi_hari_ini,
        'form': form,
        'hari_ini': now.strftime('%Y-%m-%d'),
        'batas_absen': '08:00',
    }
    return render(request, 'presensi/presensi_harian.html', konteks)


@login_required
def presensi_masuk(request):
    """
    Menangani aksi pencatatan presensi masuk.
    """
    pegawai = get_pegawai_atau_salah(request)
    if not pegawai:
        return render(request, 'presensi/no_profile.html', {'judul_halaman': 'Akses Dibatasi | Smart Presence'})

    if request.method == 'POST':
        now = timezone.localtime(timezone.now())
        today = now.date()
        current_time = now.time()

        # Validasi: Cek apakah sudah absen hari ini
        presensi_exist = Presensi.objects.filter(pegawai=pegawai, tanggal=today).exists()
        if presensi_exist:
            messages.error(request, "Anda sudah melakukan presensi masuk untuk hari ini.")
            return redirect('presensi_harian')

        form = PresensiForm(request.POST)
        keterangan = ""
        if form.is_valid():
            keterangan = form.cleaned_data.get('keterangan', '')

        # Batasan keterlambatan pukul 08:00 WIB
        batas_waktu = time(8, 0, 0)
        status = 'Hadir'
        if current_time > batas_waktu:
            status = 'Terlambat'

        # Buat data presensi baru
        Presensi.objects.create(
            pegawai=pegawai,
            tanggal=today,
            jam_masuk=current_time,
            status_kehadiran=status,
            keterangan=keterangan
        )

        messages.success(request, f"Presensi masuk berhasil dicatat pukul {current_time.strftime('%H:%M')} WIB. Status: {status}.")
    return redirect('presensi_harian')


@login_required
def presensi_pulang(request):
    """
    Menangani aksi pencatatan presensi pulang.
    """
    pegawai = get_pegawai_atau_salah(request)
    if not pegawai:
        return render(request, 'presensi/no_profile.html', {'judul_halaman': 'Akses Dibatasi | Smart Presence'})

    if request.method == 'POST':
        now = timezone.localtime(timezone.now())
        today = now.date()
        current_time = now.time()

        # Ambil data presensi hari ini
        presensi = Presensi.objects.filter(pegawai=pegawai, tanggal=today).first()

        # Validasi 1: Belum absen masuk
        if not presensi:
            messages.error(request, "Anda belum melakukan presensi masuk hari ini. Silakan presensi masuk terlebih dahulu.")
            return redirect('presensi_harian')

        # Validasi 2: Sudah absen pulang
        if presensi.jam_pulang is not None:
            messages.error(request, "Anda sudah melakukan presensi pulang untuk hari ini.")
            return redirect('presensi_harian')

        # Rekam jam pulang
        presensi.jam_pulang = current_time
        
        # Tambahkan keterangan opsional jika diisi
        keterangan_baru = request.POST.get('keterangan', '')
        if keterangan_baru:
            if presensi.keterangan:
                presensi.keterangan += f" | Pulang: {keterangan_baru}"
            else:
                presensi.keterangan = f"Pulang: {keterangan_baru}"
                
        presensi.save()
        messages.success(request, f"Presensi pulang berhasil dicatat pukul {current_time.strftime('%H:%M')} WIB. Hati-hati di jalan!")
        
    return redirect('presensi_harian')


@login_required
def presensi_riwayat(request):
    """
    Menampilkan daftar riwayat presensi pegawai.
    """
    pegawai = get_pegawai_atau_salah(request)
    if not pegawai:
        return render(request, 'presensi/no_profile.html', {'judul_halaman': 'Akses Dibatasi | Smart Presence'})

    pencarian = request.GET.get('q', '')
    riwayat = Presensi.objects.filter(pegawai=pegawai).order_by('-tanggal')

    if pencarian:
        # Filter pencarian sederhana untuk status atau keterangan
        riwayat = riwayat.filter(
            status_kehadiran__icontains=pencarian) | riwayat.filter(keterangan__icontains=pencarian
        )

    konteks = {
        'judul_halaman': 'Riwayat Kehadiran Anda | Smart Presence',
        'riwayat': riwayat,
        'pencarian': pencarian,
        'pegawai': pegawai,
    }
    return render(request, 'presensi/presensi_riwayat.html', konteks)


def get_filtered_presensi(request):
    is_admin = request.user.is_superuser or request.user.is_staff
    pegawai = get_pegawai_atau_salah(request)
    
    if is_admin:
        riwayat = Presensi.objects.all().order_by('-tanggal', '-jam_masuk')
    else:
        riwayat = Presensi.objects.filter(pegawai=pegawai).order_by('-tanggal', '-jam_masuk')
        
    tanggal_awal = request.GET.get('tanggal_awal', '')
    tanggal_akhir = request.GET.get('tanggal_akhir', '')
    pegawai_id = request.GET.get('pegawai_id', '')
    status_kehadiran = request.GET.get('status_kehadiran', '')
    
    if tanggal_awal:
        riwayat = riwayat.filter(tanggal__gte=tanggal_awal)
    if tanggal_akhir:
        riwayat = riwayat.filter(tanggal__lte=tanggal_akhir)
    if is_admin and pegawai_id:
        riwayat = riwayat.filter(pegawai_id=pegawai_id)
    if status_kehadiran:
        riwayat = riwayat.filter(status_kehadiran=status_kehadiran)
        
    filters = {
        'tanggal_awal': tanggal_awal,
        'tanggal_akhir': tanggal_akhir,
        'pegawai_id': pegawai_id,
        'status_kehadiran': status_kehadiran,
    }
    return riwayat, filters


class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(num_pages)
            super().showPage()
        super().save()

    def draw_page_number(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 9)
        self.setFillColor(colors.HexColor("#6B7280"))
        
        # Line above footer
        self.setStrokeColor(colors.HexColor("#E5E7EB"))
        self.setLineWidth(0.5)
        self.line(36, 45, 756, 45)
        
        # Footer text
        self.drawString(36, 30, "Smart Presence © 2026. Laporan Kehadiran Pegawai.")
        
        page_text = f"Halaman {self._pageNumber} dari {page_count}"
        self.drawRightString(756, 30, page_text)
        self.restoreState()


@login_required
def laporan_presensi(request):
    is_admin = request.user.is_superuser or request.user.is_staff
    pegawai = get_pegawai_atau_salah(request)
    
    if not is_admin and not pegawai:
        return render(request, 'presensi/no_profile.html', {'judul_halaman': 'Akses Dibatasi | Smart Presence'})
        
    riwayat, filters = get_filtered_presensi(request)
    
    # Calculate stats
    total_hadir = riwayat.filter(status_kehadiran='Hadir').count()
    total_terlambat = riwayat.filter(status_kehadiran='Terlambat').count()
    total_izin = riwayat.filter(status_kehadiran='Izin').count()
    total_alpa = riwayat.filter(status_kehadiran='Alpa').count()
    total_presensi = total_hadir + total_terlambat + total_izin + total_alpa
    
    if total_presensi > 0:
        persentase = round(((total_hadir + total_terlambat) / total_presensi) * 100, 1)
    else:
        persentase = 0.0
        
    stats = {
        'total_hadir': total_hadir,
        'total_terlambat': total_terlambat,
        'total_izin': total_izin,
        'total_alpa': total_alpa,
        'total_presensi': total_presensi,
        'persentase': persentase
    }
    
    daftar_pegawai = Pegawai.objects.all().order_by('nama_lengkap') if is_admin else None
    
    konteks = {
        'judul_halaman': 'Laporan Kehadiran | Smart Presence',
        'riwayat': riwayat,
        'filters': filters,
        'stats': stats,
        'is_admin': is_admin,
        'daftar_pegawai': daftar_pegawai,
    }
    return render(request, 'presensi/laporan_presensi.html', konteks)


@login_required
def export_excel(request):
    is_admin = request.user.is_superuser or request.user.is_staff
    pegawai = get_pegawai_atau_salah(request)
    
    if not is_admin and not pegawai:
        return HttpResponse("Akses ditolak: Profil pegawai tidak ditemukan.", status=403)
        
    riwayat, _ = get_filtered_presensi(request)
    
    # Prepare data for pandas DataFrame
    data = []
    for item in riwayat:
        data.append({
            'Nama Pegawai': item.pegawai.nama_lengkap,
            'Jabatan': item.pegawai.jabatan,
            'Divisi': item.pegawai.divisi,
            'Tanggal': item.tanggal.strftime('%Y-%m-%d') if item.tanggal else '-',
            'Jam Masuk': item.jam_masuk.strftime('%H:%M:%S') if item.jam_masuk else '-',
            'Jam Pulang': item.jam_pulang.strftime('%H:%M:%S') if item.jam_pulang else '-',
            'Status Kehadiran': item.status_kehadiran,
            'Keterangan': item.keterangan or '-'
        })
        
    if len(data) > 0:
        df = pd.DataFrame(data)
    else:
        df = pd.DataFrame(columns=['Nama Pegawai', 'Jabatan', 'Divisi', 'Tanggal', 'Jam Masuk', 'Jam Pulang', 'Status Kehadiran', 'Keterangan'])
    
    # Create Excel response
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="Laporan_Presensi_{}.xlsx"'.format(
        timezone.localtime(timezone.now()).strftime('%Y%m%d_%H%M%S')
    )
    
    with pd.ExcelWriter(response, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Laporan Kehadiran')
        
        workbook = writer.book
        worksheet = writer.sheets['Laporan Kehadiran']
        
        # Set column widths automatically
        for col in worksheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            worksheet.column_dimensions[col_letter].width = max(max_len + 3, 12)
            
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        
        header_fill = PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid") # Dark Navy
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        center_align = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style='thin', color='DDDDDD'),
            right=Side(style='thin', color='DDDDDD'),
            top=Side(style='thin', color='DDDDDD'),
            bottom=Side(style='thin', color='DDDDDD')
        )
        
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            
        for row in worksheet.iter_rows(min_row=2, max_row=len(data) + 1, min_col=1, max_col=8):
            for cell in row:
                cell.border = thin_border
                
    return response


@login_required
def export_pdf(request):
    is_admin = request.user.is_superuser or request.user.is_staff
    pegawai = get_pegawai_atau_salah(request)
    
    if not is_admin and not pegawai:
        return HttpResponse("Akses ditolak: Profil pegawai tidak ditemukan.", status=403)
        
    riwayat, filters = get_filtered_presensi(request)
    
    # Calculate stats
    total_hadir = riwayat.filter(status_kehadiran='Hadir').count()
    total_terlambat = riwayat.filter(status_kehadiran='Terlambat').count()
    total_izin = riwayat.filter(status_kehadiran='Izin').count()
    total_alpa = riwayat.filter(status_kehadiran='Alpa').count()
    total_presensi = total_hadir + total_terlambat + total_izin + total_alpa
    persentase = round(((total_hadir + total_terlambat) / total_presensi) * 100, 1) if total_presensi > 0 else 0.0
    
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="Laporan_Presensi_{}.pdf"'.format(
        timezone.localtime(timezone.now()).strftime('%Y%m%d_%H%M%S')
    )
    
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=54
    )
    
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=24,
        leading=28,
        textColor=colors.HexColor('#1E1B4B')
    )
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#4B5563')
    )
    heading_style = ParagraphStyle(
        'SectionHeading',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#1E1B4B')
    )
    table_header_style = ParagraphStyle(
        'TableHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        leading=11,
        textColor=colors.white,
        alignment=1
    )
    table_cell_style = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=11,
        textColor=colors.HexColor('#1F2937')
    )
    table_cell_center = ParagraphStyle(
        'TableCellCenter',
        parent=table_cell_style,
        alignment=1
    )
    meta_label_style = ParagraphStyle(
        'MetaLabel',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#4B5563')
    )
    meta_val_style = ParagraphStyle(
        'MetaValue',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#1F2937')
    )
    
    status_hadir = ParagraphStyle('StatHadir', parent=table_cell_center, fontName='Helvetica-Bold', textColor=colors.HexColor('#10B981'))
    status_terlambat = ParagraphStyle('StatTerlambat', parent=table_cell_center, fontName='Helvetica-Bold', textColor=colors.HexColor('#F59E0B'))
    status_izin = ParagraphStyle('StatIzin', parent=table_cell_center, fontName='Helvetica-Bold', textColor=colors.HexColor('#3B82F6'))
    status_alpa = ParagraphStyle('StatAlpa', parent=table_cell_center, fontName='Helvetica-Bold', textColor=colors.HexColor('#EF4444'))
    
    story = []
    
    # Header
    header_data = [
        [
            Paragraph("SMART <font color='#06B6D4'>PRESENCE</font>", title_style),
            Paragraph(f"TANGGAL CETAK: {timezone.localtime(timezone.now()).strftime('%d %B %Y, %H:%M')} WIB", ParagraphStyle('RightMeta', parent=meta_val_style, alignment=2))
        ],
        [
            Paragraph("SISTEM MONITORING & LAPORAN KEHADIRAN PEGAWAI PINTAR", subtitle_style),
            ""
        ]
    ]
    header_table = Table(header_data, colWidths=[400, 320])
    header_table.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 0),
        ('TOPPADDING', (0,0), (-1,-1), 0),
    ]))
    story.append(header_table)
    story.append(Spacer(1, 10))
    
    # Divider
    bar_data = [['']]
    bar_table = Table(bar_data, colWidths=[720], rowHeights=[3])
    bar_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#1E1B4B')),
        ('TOPPADDING', (0,0), (-1,-1), 0),
        ('BOTTOMPADDING', (0,0), (-1,-1), 0),
    ]))
    story.append(bar_table)
    story.append(Spacer(1, 15))
    
    # Filter metadata table
    meta_info = []
    meta_info.append([Paragraph("Saringan Periode:", meta_label_style), Paragraph(f"{filters['tanggal_awal'] or 'Mulai Awal'} s/d {filters['tanggal_akhir'] or 'Hari Ini'}", meta_val_style)])
    if is_admin:
        peg_nama = "Semua Pegawai"
        if filters['pegawai_id']:
            p_obj = Pegawai.objects.filter(id=filters['pegawai_id']).first()
            if p_obj:
                peg_nama = p_obj.nama_lengkap
        meta_info.append([Paragraph("Saringan Pegawai:", meta_label_style), Paragraph(peg_nama, meta_val_style)])
    else:
        meta_info.append([Paragraph("Pegawai:", meta_label_style), Paragraph(pegawai.nama_lengkap, meta_val_style)])
        
    meta_info.append([Paragraph("Saringan Status:", meta_label_style), Paragraph(filters['status_kehadiran'] or "Semua Status", meta_val_style)])
    
    meta_table = Table(meta_info, colWidths=[120, 220])
    meta_table.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('LEFTPADDING', (0,0), (-1,-1), 0),
    ]))
    
    # Statistics summary
    stats_data = [
        [
            Paragraph("<b>Hadir</b>", meta_label_style),
            Paragraph("<b>Terlambat</b>", meta_label_style),
            Paragraph("<b>Izin</b>", meta_label_style),
            Paragraph("<b>Alpa</b>", meta_label_style),
            Paragraph("<b>Persentase</b>", meta_label_style),
        ],
        [
            Paragraph(f"<font color='#10B981'><b>{total_hadir}</b></font> Hari", meta_val_style),
            Paragraph(f"<font color='#F59E0B'><b>{total_terlambat}</b></font> Kali", meta_val_style),
            Paragraph(f"<font color='#3B82F6'><b>{total_izin}</b></font> Hari", meta_val_style),
            Paragraph(f"<font color='#EF4444'><b>{total_alpa}</b></font> Hari", meta_val_style),
            Paragraph(f"<b>{persentase}%</b>", meta_val_style),
        ]
    ]
    stats_table = Table(stats_data, colWidths=[68, 68, 68, 68, 80])
    stats_table.setStyle(TableStyle([
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#F3F4F6')),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E5E7EB')),
        ('TOPPADDING', (0,0), (-1,-1), 6),
        ('BOTTOMPADDING', (0,0), (-1,-1), 6),
    ]))
    
    combined_data = [[meta_table, stats_table]]
    combined_table = Table(combined_data, colWidths=[360, 360])
    combined_table.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('LEFTPADDING', (0,0), (-1,-1), 0),
        ('RIGHTPADDING', (0,0), (-1,-1), 0),
        ('BOTTOMPADDING', (0,0), (-1,-1), 15),
    ]))
    story.append(combined_table)
    
    # Table data
    table_data = [[
        Paragraph("No.", table_header_style),
        Paragraph("Nama Pegawai", table_header_style),
        Paragraph("Jabatan / Divisi", table_header_style),
        Paragraph("Tanggal", table_header_style),
        Paragraph("Jam Masuk", table_header_style),
        Paragraph("Jam Pulang", table_header_style),
        Paragraph("Status", table_header_style),
        Paragraph("Keterangan", table_header_style),
    ]]
    
    for idx, item in enumerate(riwayat, 1):
        st_val = item.status_kehadiran
        if st_val == 'Hadir':
            st_p = Paragraph(st_val, status_hadir)
        elif st_val == 'Terlambat':
            st_p = Paragraph(st_val, status_terlambat)
        elif st_val == 'Izin':
            st_p = Paragraph(st_val, status_izin)
        else:
            st_p = Paragraph(st_val, status_alpa)
            
        table_data.append([
            Paragraph(str(idx), table_cell_center),
            Paragraph(item.pegawai.nama_lengkap, table_cell_style),
            Paragraph(f"{item.pegawai.jabatan}<br/><font color='#6B7280'>{item.pegawai.divisi}</font>", table_cell_style),
            Paragraph(item.tanggal.strftime('%d-%m-%Y') if item.tanggal else '-', table_cell_center),
            Paragraph(item.jam_masuk.strftime('%H:%M:%S') if item.jam_masuk else '-', table_cell_center),
            Paragraph(item.jam_pulang.strftime('%H:%M:%S') if item.jam_pulang else '-', table_cell_center),
            st_p,
            Paragraph(item.keterangan or '-', table_cell_style),
        ])
        
    if len(table_data) == 1:
        table_data.append([
            Paragraph("-", table_cell_center),
            Paragraph("Tidak ada data presensi yang sesuai dengan kriteria saringan.", table_cell_style),
            "", "", "", "", "", ""
        ])
        
    data_table = Table(table_data, colWidths=[30, 130, 120, 80, 70, 70, 80, 140])
    
    t_styles = [
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1E1B4B')),
        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E5E7EB')),
        ('TOPPADDING', (0,0), (-1,-1), 6),
        ('BOTTOMPADDING', (0,0), (-1,-1), 6),
    ]
    
    if len(table_data) > 2 or (len(table_data) == 2 and table_data[1][1] != "Tidak ada data presensi yang sesuai dengan kriteria saringan."):
        for row_idx in range(1, len(table_data)):
            if row_idx % 2 == 0:
                t_styles.append(('BACKGROUND', (0, row_idx), (-1, row_idx), colors.HexColor('#F9FAFB')))
    else:
        t_styles.append(('SPAN', (1, 1), (7, 1)))
        t_styles.append(('ALIGN', (1, 1), (7, 1), 'CENTER'))
        
    data_table.setStyle(TableStyle(t_styles))
    story.append(data_table)
    
    doc.build(story, canvasmaker=NumberedCanvas)
    
    pdf_content = buffer.getvalue()
    buffer.close()
    
    response.write(pdf_content)
    return response


@login_required
def ajukan_izin(request):
    """
    Halaman bagi pegawai untuk mengajukan izin, sakit, atau cuti.
    """
    pegawai = get_pegawai_atau_salah(request)
    if not pegawai:
        return render(request, 'presensi/no_profile.html', {'judul_halaman': 'Akses Dibatasi | Smart Presence'})

    if request.method == 'POST':
        form = PengajuanIzinForm(request.POST, request.FILES)
        if form.is_valid():
            pengajuan = form.save(commit=False)
            pengajuan.pegawai = pegawai

            # Validasi penanggalan
            if pengajuan.tanggal_mulai > pengajuan.tanggal_selesai:
                messages.error(request, "Tanggal selesai tidak boleh mendahului tanggal mulai.")
            else:
                pengajuan.save()
                messages.success(request, "Pengajuan izin berhasil dikirimkan dan sedang menunggu peninjauan admin.")
                return redirect('riwayat_izin')
    else:
        form = PengajuanIzinForm()

    konteks = {
        'judul_halaman': 'Ajukan Izin | Smart Presence',
        'form': form,
        'pegawai': pegawai,
    }
    return render(request, 'presensi/ajukan_izin.html', konteks)


@login_required
def riwayat_izin(request):
    """
    Menampilkan riwayat pengajuan izin personal pegawai.
    """
    is_admin = request.user.is_superuser or request.user.is_staff
    pegawai = get_pegawai_atau_salah(request)
    
    if not pegawai:
        if is_admin:
            return redirect('admin_izin')
        return render(request, 'presensi/no_profile.html', {'judul_halaman': 'Akses Dibatasi | Smart Presence'})

    riwayat = PengajuanIzin.objects.filter(pegawai=pegawai).order_by('-dibuat_pada')

    konteks = {
        'judul_halaman': 'Riwayat Izin Anda | Smart Presence',
        'riwayat': riwayat,
        'pegawai': pegawai,
    }
    return render(request, 'presensi/riwayat_izin.html', konteks)


@login_required
def admin_izin(request):
    """
    Halaman manajemen kelola pengajuan izin bagi admin/superuser.
    """
    if not (request.user.is_superuser or request.user.is_staff):
        messages.error(request, "Akses ditolak: Area khusus administrator.")
        return redirect('dashboard')

    riwayat = PengajuanIzin.objects.all().order_by('-dibuat_pada')

    # Filter pencarian
    status_filter = request.GET.get('status_persetujuan', '')
    pegawai_filter = request.GET.get('pegawai_id', '')

    if status_filter:
        riwayat = riwayat.filter(status_persetujuan=status_filter)
    if pegawai_filter:
        riwayat = riwayat.filter(pegawai_id=pegawai_filter)

    daftar_pegawai = Pegawai.objects.all().order_by('nama_lengkap')

    konteks = {
        'judul_halaman': 'Persetujuan Izin | Smart Presence',
        'riwayat': riwayat,
        'daftar_pegawai': daftar_pegawai,
        'filters': {
            'status_persetujuan': status_filter,
            'pegawai_id': pegawai_filter,
        }
    }
    return render(request, 'presensi/admin_izin.html', konteks)


@login_required
def setujui_izin(request, izin_id):
    """
    Menyetujui pengajuan izin pegawai dan mengintegrasikannya ke tabel Presensi harian.
    """
    if not (request.user.is_superuser or request.user.is_staff):
        messages.error(request, "Akses ditolak: Area khusus administrator.")
        return redirect('dashboard')

    pengajuan = PengajuanIzin.objects.filter(id=izin_id).first()
    if not pengajuan:
        messages.error(request, "Pengajuan izin tidak ditemukan.")
        return redirect('admin_izin')

    if request.method == 'POST':
        catatan = request.POST.get('catatan_admin', '')
        pengajuan.status_persetujuan = 'Disetujui'
        pengajuan.catatan_admin = catatan
        pengajuan.save()

        # Otomatis buat data di tabel Presensi harian untuk seluruh tanggal pengajuan
        import datetime
        curr_date = pengajuan.tanggal_mulai
        end_date = pengajuan.tanggal_selesai

        while curr_date <= end_date:
            presensi_exist = Presensi.objects.filter(pegawai=pengajuan.pegawai, tanggal=curr_date).first()
            keterangan_text = f"Izin disetujui ({pengajuan.jenis_izin}): {pengajuan.alasan}"
            
            if presensi_exist:
                presensi_exist.status_kehadiran = 'Izin'
                if presensi_exist.keterangan:
                    if keterangan_text not in presensi_exist.keterangan:
                        presensi_exist.keterangan += f" | {keterangan_text}"
                else:
                    presensi_exist.keterangan = keterangan_text
                presensi_exist.save()
            else:
                Presensi.objects.create(
                    pegawai=pengajuan.pegawai,
                    tanggal=curr_date,
                    status_kehadiran='Izin',
                    keterangan=keterangan_text
                )
            curr_date += datetime.timedelta(days=1)

        messages.success(request, f"Pengajuan izin {pengajuan.pegawai.nama_lengkap} berhasil DISETUJUI dan data presensi telah diperbarui.")
    
    return redirect('admin_izin')


@login_required
def tolak_izin(request, izin_id):
    """
    Menolak pengajuan izin pegawai dengan menyertakan alasan penolakan.
    """
    if not (request.user.is_superuser or request.user.is_staff):
        messages.error(request, "Akses ditolak: Area khusus administrator.")
        return redirect('dashboard')

    pengajuan = PengajuanIzin.objects.filter(id=izin_id).first()
    if not pengajuan:
        messages.error(request, "Pengajuan izin tidak ditemukan.")
        return redirect('admin_izin')

    if request.method == 'POST':
        catatan = request.POST.get('catatan_admin', '')
        pengajuan.status_persetujuan = 'Ditolak'
        pengajuan.catatan_admin = catatan
        pengajuan.save()
        
        messages.warning(request, f"Pengajuan izin {pengajuan.pegawai.nama_lengkap} telah DITOLAK.")
        
    return redirect('admin_izin')
